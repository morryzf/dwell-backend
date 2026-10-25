import asyncio
import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from fastapi.testclient import TestClient

from app import auth, db, file_make
from app.main import FILE_TOOL_NAME, _chat_tools, _heartbeat_tool_allowed, _run_chat_tool, app

SAMPLE = """# 十月复习计划

这是一份**心理测量学**的复习计划。

## 第一周

- 信度
  - 重测信度
- 效度

1. 先看课本
2. 再做习题

> 慢慢来。

| 周次 | 内容 |
|---|---|
| 1 | 信度 |

---

```
alpha = k/(k-1)
```
"""


class MakeFileTest(unittest.TestCase):
    """文件名带扩展名决定格式；文字类原样存，Word / PDF 从 Markdown 排出来。"""

    def test_text_file_is_stored_as_is(self):
        name, mime, data = file_make.make_file("笔记.md", "# 标题\n内容")
        self.assertEqual((name, mime), ("笔记.md", "text/markdown"))
        self.assertEqual(data.decode("utf-8"), "# 标题\n内容")

    def test_csv_gets_bom_for_excel(self):
        _, _, data = file_make.make_file("清单.csv", "科目,时长\n信度,6\n")
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))

    def test_docx_has_headings_lists_and_table(self):
        import docx

        _, mime, data = file_make.make_file("计划.docx", SAMPLE)
        self.assertIn("wordprocessingml", mime)
        document = docx.Document(io.BytesIO(data))
        styles = [p.style.name for p in document.paragraphs]
        self.assertIn("Heading 1", styles)
        self.assertIn("List Bullet", styles)
        self.assertIn("List Bullet 2", styles)
        self.assertIn("List Number", styles)
        self.assertIn("Quote", styles)
        self.assertEqual(len(document.tables), 1)
        bold = [run.text for p in document.paragraphs for run in p.runs if run.bold]
        self.assertIn("心理测量学", bold)

    def test_pdf_embeds_chinese_text(self):
        if not file_make.find_cjk_font()[0]:
            self.skipTest("这台机器没有中文字体")
        from pypdf import PdfReader

        _, mime, data = file_make.make_file("计划.pdf", SAMPLE)
        self.assertEqual(mime, "application/pdf")
        text = PdfReader(io.BytesIO(data)).pages[0].extract_text()
        self.assertIn("十月复习计划", text.replace("​", ""))

    def test_pdf_without_font_says_so(self):
        original = file_make.find_cjk_font
        file_make.find_cjk_font = lambda: ("", "")
        try:
            with self.assertRaisesRegex(ValueError, "中文字体"):
                file_make.make_file("计划.pdf", "内容")
        finally:
            file_make.find_cjk_font = original

    def test_bad_names_and_content(self):
        for name in ("没有扩展名", "病毒.exe"):
            with self.assertRaises(ValueError):
                file_make.make_file(name, "内容")
        with self.assertRaises(ValueError):
            file_make.make_file("空.txt", "   ")
        self.assertEqual(file_make.clean_name("../../etc/密码.txt"), "密码.txt")

    def test_chinese_gets_break_points_but_not_before_closing_punctuation(self):
        text = file_make.breakable("中文，好。abc")
        self.assertNotIn("​，", text)
        self.assertNotIn("​。", text)
        self.assertIn("中​文", text)
        self.assertNotIn("a​b", text)


class FileToolFlowTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.previous_path = db.DB_PATH
        db.DB_PATH = str(Path(cls.tmp.name) / "files.db")
        db.init_db()
        app.dependency_overrides[auth.require_auth] = lambda: None
        cls.client = TestClient(app)
        cls.chat = db.chat_add("文件")

    @classmethod
    def tearDownClass(cls):
        app.dependency_overrides.pop(auth.require_auth, None)
        db.DB_PATH = cls.previous_path
        cls.tmp.cleanup()

    def test_tool_is_offered_in_every_chat_but_not_to_heartbeat(self):
        tools, tool_map = asyncio.run(_chat_tools(self.chat["id"], list_tools=lambda server: []))
        names = [tool["function"]["name"] for tool in tools]
        self.assertIn(FILE_TOOL_NAME, names)
        tool = next(tool for tool in tools if tool["function"]["name"] == FILE_TOOL_NAME)
        self.assertFalse(_heartbeat_tool_allowed(tool, tool_map[FILE_TOOL_NAME]))

    def test_created_file_downloads_with_its_name(self):
        result, is_error = asyncio.run(_run_chat_tool(
            self.chat["id"], FILE_TOOL_NAME, "builtin:file", {"filename": "笔记.txt", "content": "你好"}))
        self.assertFalse(is_error)
        file = json.loads(result)["file"]
        response = self.client.get("/api/files/" + file["id"])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content.decode("utf-8"), "你好")
        self.assertIn("charset=utf-8", response.headers["content-type"])
        self.assertIn("inline", response.headers["content-disposition"])
        self.assertIn("filename*=UTF-8''%E7%AC%94%E8%AE%B0.txt", response.headers["content-disposition"])
        saved = self.client.get("/api/files/" + file["id"] + "?download=1")
        self.assertIn("attachment", saved.headers["content-disposition"])

    def test_html_is_sandboxed(self):
        result, _ = asyncio.run(_run_chat_tool(
            self.chat["id"], FILE_TOOL_NAME, "builtin:file",
            {"filename": "页.html", "content": "<script>alert(1)</script>"}))
        file = json.loads(result)["file"]
        response = self.client.get("/api/files/" + file["id"])
        self.assertEqual(response.headers.get("content-security-policy"), "sandbox")

    def test_tool_error_goes_back_to_the_model(self):
        result, is_error = asyncio.run(_run_chat_tool(
            self.chat["id"], FILE_TOOL_NAME, "builtin:file", {"filename": "x.exe", "content": "a"}))
        self.assertTrue(is_error)
        self.assertIn("不支持", result)

    def test_missing_file_is_404(self):
        self.assertEqual(self.client.get("/api/files/nope").status_code, 404)


class FileCardStaticTest(unittest.TestCase):
    def test_page_draws_a_card_for_created_files(self):
        html = (Path(__file__).resolve().parents[1] / "static" / "index.html").read_text(encoding="utf-8")
        self.assertIn("if (t.name === 'DwellFileCreate' && !isError) addGeneratedFileCard(t);", html)
        self.assertIn("save.href = 'api/files/' + encodeURIComponent(file.id) + '?download=1';", html)


if __name__ == "__main__":
    unittest.main()
